import os
import fitz  # PyMuPDF
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pydantic import BaseModel, Field
from google import genai
from google.genai import types

# -----------------------------------------------------------------------------
# 1. Setup & Initialization
# -----------------------------------------------------------------------------
api_key = os.getenv("GEMINI_API_KEY")
if not api_key:
    raise ValueError("Please set GEMINI_API_KEY environment variable.")

client = genai.Client(api_key=api_key)

def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip('#')
    return RGBColor(*(int(hex_str[i:i+2], 16) for i in (0, 2, 4)))

# -----------------------------------------------------------------------------
# 2. Extract PDF Visuals
# -----------------------------------------------------------------------------
def extract_pdf_images(pdf_path: str, output_folder: str = "extracted_assets") -> list[str]:
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)

    extracted_files = []
    doc = fitz.open(pdf_path)
    image_count = 0

    for page_index in range(len(doc)):
        page = doc[page_index]
        for img_index, img_info in enumerate(page.get_images(full=True)):
            xref = img_info[0]
            base_image = doc.extract_image(xref)
            image_bytes = base_image["image"]
            image_ext = base_image["ext"]

            if len(image_bytes) < 15000:  # Skip icons/micro-graphics
                continue

            img_filename = os.path.join(output_folder, f"img_p{page_index + 1}_{img_index + 1}.{image_ext}")
            with open(img_filename, "wb") as f:
                f.write(image_bytes)

            extracted_files.append(img_filename)
            image_count += 1
            if image_count >= 15:
                break
        if image_count >= 15:
            break

    doc.close()
    return extracted_files

# -----------------------------------------------------------------------------
# 3. Pydantic Extraction Schemas
# -----------------------------------------------------------------------------
class MetricItem(BaseModel):
    number: str = Field(description="Stat value, e.g. '160+' or '16,395'")
    label: str = Field(description="Short label describing the metric")

class SlideContent(BaseModel):
    slide_type: str = Field(description="'chapter_cover', 'split_content', or 'metrics_grid'")
    title: str = Field(description="Main slide title")
    subtitle: str = Field(description="Optional subtitle or context", default="")
    chapter_number: str = Field(description="Chapter number if cover", default="")
    bullet_points: list[str] = Field(description="Executive bullet points", default_factory=list)
    metrics: list[MetricItem] = Field(description="Key metrics if grid", default_factory=list)
    image_index: int = Field(description="Index of extracted image asset to show (-1 if none)", default=-1)

class PresentationSchema(BaseModel):
    company_name: str = Field(description="Company or Client name")
    project_title: str = Field(description="Project or Proposal title")
    primary_color: str = Field(description="Primary brand hex code", default="#0B192C")
    accent_color: str = Field(description="Accent brand hex code", default="#8B1538")
    slides: list[SlideContent] = Field(description="Ordered list of slides")

# -----------------------------------------------------------------------------
# 4. Modern PowerPoint Layout Engine
# -----------------------------------------------------------------------------
class ModernDeckBuilder:
    def __init__(self, data: PresentationSchema, images: list[str]):
        self.data = data
        self.images = images
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)

        # Executive Modern Palette
        self.c_primary = hex_to_rgb(self.data.primary_color)
        self.c_accent = hex_to_rgb(self.data.accent_color)
        self.c_bg = hex_to_rgb("#F8FAFC")         # Crisp Off-White/Slate
        self.c_card_bg = hex_to_rgb("#FFFFFF")    # Clean White Cards
        self.c_card_border = hex_to_rgb("#E2E8F0")# Subtle Border
        self.c_text_dark = hex_to_rgb("#0F172A")  # Slate Dark
        self.c_text_muted = hex_to_rgb("#64748B") # Muted Subtitles

        self.f_head = "Segoe UI"
        self.f_body = "Calibri"

    def apply_bg(self, slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(self, slide, title, subtitle):
        # Vertical Pill Accent Indicator
        pill = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.5), Inches(0.1), Inches(0.75))
        pill.fill.solid()
        pill.fill.fore_color.rgb = self.c_accent
        pill.line.fill.background()

        # Title & Subtitle Box
        tb = slide.shapes.add_textbox(Inches(1.05), Inches(0.4), Inches(11.4), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = self.f_head
        p0.font.size = Pt(22)
        p0.font.bold = True
        p0.font.color.rgb = self.c_primary

        if subtitle:
            p1 = tf.add_paragraph()
            p1.text = subtitle
            p1.font.name = self.f_body
            p1.font.size = Pt(13)
            p1.font.color.rgb = self.c_text_muted

    def add_footer(self, slide, num):
        # Footer Divider Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.85), Inches(11.733), Inches(0.01))
        line.fill.solid()
        line.fill.fore_color.rgb = self.c_card_border
        line.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = f"{self.data.company_name}   |   {self.data.project_title}"
        p.font.name = self.f_body
        p.font.size = Pt(10)
        p.font.color.rgb = self.c_text_muted

        p_num = tf.add_paragraph()
        p_num.text = f"{num}"
        p_num.font.name = self.f_body
        p_num.font.size = Pt(10)
        p_num.font.color.rgb = self.c_text_muted
        p_num.alignment = PP_ALIGN.RIGHT

    def add_chapter_cover(self, slide_data: SlideContent, num: int):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.apply_bg(slide, self.c_primary)

        # Left Accent Block
        block = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.3), Inches(7.5))
        block.fill.solid()
        block.fill.fore_color.rgb = self.c_accent
        block.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.5), Inches(4.0))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = f"CHAPTER {slide_data.chapter_number.zfill(2)}"
        p0.font.name = self.f_head
        p0.font.size = Pt(18)
        p0.font.bold = True
        p0.font.color.rgb = self.c_accent
        p0.space_after = Pt(12)

        p1 = tf.add_paragraph()
        p1.text = slide_data.title.upper()
        p1.font.name = self.f_head
        p1.font.size = Pt(36)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(255, 255, 255)
        p1.space_after = Pt(12)

        if slide_data.subtitle:
            p2 = tf.add_paragraph()
            p2.text = slide_data.subtitle
            p2.font.name = self.f_body
            p2.font.size = Pt(16)
            p2.font.color.rgb = RGBColor(203, 213, 225)

    def add_split_slide(self, slide_data: SlideContent, num: int):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.apply_bg(slide, self.c_bg)
        self.add_header(slide, slide_data.title, slide_data.subtitle)

        has_image = 0 <= slide_data.image_index < len(self.images)
        card_w = Inches(6.0) if has_image else Inches(11.733)

        # Main Text Card Container
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), card_w, Inches(5.1))
        card.fill.solid()
        card.fill.fore_color.rgb = self.c_card_bg
        card.line.color.rgb = self.c_card_border

        tx = slide.shapes.add_textbox(Inches(1.1), Inches(1.7), card_w - Inches(0.6), Inches(4.7))
        tf = tx.text_frame
        tf.word_wrap = True

        for i, b in enumerate(slide_data.bullet_points):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = f"•  {b}"
            p.font.name = self.f_body
            p.font.size = Pt(14)
            p.font.color.rgb = self.c_text_dark
            p.space_after = Pt(14)

        # Right Side Image Container
        if has_image:
            img_path = self.images[slide_data.image_index]
            if os.path.exists(img_path):
                slide.shapes.add_picture(img_path, Inches(7.1), Inches(1.5), width=Inches(5.433))

        self.add_footer(slide, num)

    def add_metrics_grid(self, slide_data: SlideContent, num: int):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.apply_bg(slide, self.c_bg)
        self.add_header(slide, slide_data.title, slide_data.subtitle)

        metrics = slide_data.metrics[:4]
        if metrics:
            n_cards = len(metrics)
            w = Inches(11.733 / n_cards - 0.2)

            for i, m in enumerate(metrics):
                left = Inches(0.8) + i * (w + Inches(0.2))

                # Stat Card
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), w, Inches(4.7))
                card.fill.solid()
                card.fill.fore_color.rgb = self.c_card_bg
                card.line.color.rgb = self.c_card_border

                # Big Stat Callout
                tb_num = slide.shapes.add_textbox(left, Inches(2.3), w, Inches(1.0))
                p_n = tb_num.text_frame.paragraphs[0]
                p_n.text = str(m.number)
                p_n.font.name = self.f_head
                p_n.font.size = Pt(36)
                p_n.font.bold = True
                p_n.font.color.rgb = self.c_accent
                p_n.alignment = PP_ALIGN.CENTER

                # Stat Label
                tb_lbl = slide.shapes.add_textbox(left + Inches(0.1), Inches(3.5), w - Inches(0.2), Inches(2.0))
                tf_l = tb_lbl.text_frame
                tf_l.word_wrap = True
                p_l = tf_l.paragraphs[0]
                p_l.text = str(m.label)
                p_l.font.name = self.f_body
                p_l.font.size = Pt(14)
                p_l.font.color.rgb = self.c_text_dark
                p_l.alignment = PP_ALIGN.CENTER

        self.add_footer(slide, num)

    def build(self, filename="Modernized_Presentation.pptx"):
        for i, slide in enumerate(self.data.slides):
            if slide.slide_type == "chapter_cover":
                self.add_chapter_cover(slide, i + 1)
            elif slide.slide_type == "metrics_grid":
                self.add_metrics_grid(slide, i + 1)
            else:
                self.add_split_slide(slide, i + 1)

        self.prs.save(filename)
        print(f"🎉 Modernized PowerPoint successfully built: '{filename}'")

# -----------------------------------------------------------------------------
# 5. Execution Pipeline
# -----------------------------------------------------------------------------
def run_pipeline(pdf_path: str):
    images = extract_pdf_images(pdf_path)
    file_ref = client.files.upload(file=pdf_path)

    prompt = f"""
    Analyze the attached PDF proposal and structure it into a modern executive slide deck blueprint.
    
    1. Extract company name and project title.
    2. Divide into logical chapters ('chapter_cover').
    3. Use 'split_content' for descriptive slides, creating 3-4 concise bullet points.
    4. Use 'metrics_grid' for slides rich in numerical statistics.
    5. Assign image indices (0 to {max(len(images)-1, 0)}) to relevant slides where visual diagrams apply.
    """

    res = client.models.generate_content(
        model="gemini-3.5-flash",
        contents=[file_ref, prompt],
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=PresentationSchema,
        ),
    )

    extracted = PresentationSchema.model_validate_json(res.text)
    builder = ModernDeckBuilder(extracted, images)
    builder.build(filename=f"{extracted.company_name.replace(' ', '_')}_Modern_Deck.pptx")

if __name__ == "__main__":
    target_pdf = "QDB Bluu Proposal v10 (low-res).pdf"
    if os.path.exists(target_pdf):
        run_pipeline(target_pdf)
    else:
        print(f"File not found: {target_pdf}")
